#!/usr/bin/env python3
"""Build a private editable photo-reference review through the existing compiler.

The host model inspects pixels and records the rectified layout. An explicitly
estimated review-only route requires model provenance rather than invented
per-coordinate or per-font user approval. This adapter
does no OCR, perspective correction or model/API call. It is the project-private
create-template owner handoff, not a registered-template/new-deck gate bypass.

Usage: python3 scripts/photo_reconstruct.py PLAN.json --check
Examples: python3 scripts/photo_reconstruct.py PLAN.json --workspace-parent projects/private --presentations-skill-dir <verified-shared-Presentations-skill>
Dependencies: existing prepared Pillow/fontTools/python-pptx and shared Presentations runtime
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
from xml.etree import ElementTree as ET

from console_encoding import configure_utf8_stdio
from korean_text_fit import measure_text_fit
from private_font_cache import inspect_font
from project_manager import ProjectManager
from reference_intake import IMAGE_TYPES, bind_photo_observations, inspect_reference
from template_gate import make_exempt_record, write_project_gate

configure_utf8_stdio()
SCRIPTS = Path(__file__).resolve().parent
SKILL = SCRIPTS.parent
REPO = SKILL.parents[2]
SVG = 'http://www.w3.org/2000/svg'
ET.register_namespace('', SVG)
CANVAS = [1280, 720]
SHA = re.compile(r'[a-f0-9]{64}')
ID = re.compile(r'[A-Za-z][A-Za-z0-9_-]{0,79}')


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def verify_fitted_font_files(audit: dict) -> dict[str, str]:
    """Keep the font bytes used for glyph fitting bound to later stages."""
    files, faces = audit['font_files'], audit['font_faces']
    if not files or len(files) != len(faces):
        raise ValueError('Fitted font file/hash inventory is incomplete')
    expected = {str(Path(path).resolve()): face['sha256'] for path, face in zip(files, faces)}
    if len(expected) != len(files) or any(not isinstance(sha, str) or not SHA.fullmatch(sha) for sha in expected.values()):
        raise ValueError('Fitted font file/hash inventory is invalid')
    if any(digest(Path(path)) != sha for path, sha in expected.items()):
        raise ValueError('Font bytes changed after text-fit validation')
    return expected


def verify_render_font_binding(audit: dict, manifest: dict) -> dict:
    expected = verify_fitted_font_files(audit)
    if not isinstance(manifest, dict) or manifest.get('registered_fonts') != expected:
        raise ValueError('Render font registration must match every fitted font path and SHA256')
    return {'schema_version': 1, 'fit_render_font_hashes_verified': True,
            'selected_font_files_sha256': expected, 'source_font_identity_verified': False}


def label(value: object, name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{name} requires a nonempty provenance/confirmation reference')
    return value


def number(value: object, name: str, *, positive: bool = True) -> float:
    if type(value) not in (int, float) or not math.isfinite(value) or (value <= 0 if positive else value < 0):
        raise ValueError(f'{name} requires a finite {"positive" if positive else "nonnegative"} number')
    return value


def box(value: object, size: list, name: str) -> list:
    if not isinstance(value, list) or len(value) != 4:
        raise ValueError(f'{name} requires [x,y,width,height] in px')
    x, y, w, h = value
    number(x, name, positive=False); number(y, name, positive=False)
    number(w, name); number(h, name)
    if x + w > size[0] or y + h > size[1]:
        raise ValueError(f'{name} exceeds its coordinate space')
    return value


def color(value: object) -> str:
    if not isinstance(value, str) or not re.fullmatch(r'#[0-9A-Fa-f]{6}', value):
        raise ValueError('Explicit six-digit hex color required')
    return value


def local_file(base: Path, value: object) -> Path:
    raw = Path(label(value, 'local file'))
    return (raw if raw.is_absolute() else base / raw).resolve()


def validate_plan(plan: dict, base: Path) -> dict:
    """Validate before any output write; retain both photo and slide spaces."""
    if not isinstance(plan, dict): raise ValueError('Reconstruction plan must be an object')
    if type(plan.get('schema_version')) is not int or plan['schema_version'] != 1:
        raise ValueError('Reconstruction plan requires schema_version 1')
    if plan.get('purpose_choice') != 'copy_original' or plan.get('purpose_confirmed') is not True:
        raise ValueError('Actual current-user copy_original confirmation required')
    label(plan.get('purpose_confirmation_ref'), 'purpose')
    if 'review_only' in plan and type(plan['review_only']) is not bool:
        raise ValueError('review_only must be a Boolean')
    estimated_review = plan.get('review_only') is True
    if estimated_review:
        if plan.get('layout_basis') != 'visual_estimate':
            raise ValueError('Review-only layout must retain visual_estimate basis')
        label(plan.get('layout_review_ref'), 'model layout review')
        if ('layout_confirmed' in plan and plan['layout_confirmed'] is not False) or 'layout_confirmation_ref' in plan:
            raise ValueError('Model estimate cannot be labeled as confirmed layout approval')
    else:
        if plan.get('layout_confirmed') is not True:
            raise ValueError('Current model reconstruction layout must be confirmed')
        label(plan.get('layout_confirmation_ref'), 'layout')
        if 'layout_basis' in plan or 'layout_review_ref' in plan:
            raise ValueError('Estimated layout requires explicit review_only route')
    if plan.get('canvas_px') != CANVAS:
        raise ValueError('Minimal adapter supports only explicit 1280x720 (16:9) canvas')
    name = plan.get('project_name')
    if not isinstance(name, str) or not ID.fullmatch(name):
        raise ValueError('Safe project_name required; no path characters')
    choice = plan.get('font_choice', {})
    if not isinstance(choice, dict): raise ValueError('font_choice must be an object')
    family_name = label(choice.get('family'), 'selected font family')
    if choice.get('source_font_identity_verified') is not False:
        raise ValueError('Photo source font identity must remain unverified')
    if estimated_review:
        if choice.get('basis') != 'visual_estimate' or 'confirmation_ref' in choice:
            raise ValueError('Model font estimate needs review provenance, not user confirmation')
        label(choice.get('review_ref'), 'model font selection')
    else:
        if choice.get('basis') != 'design': raise ValueError('Confirmed path requires design font choice')
        label(choice.get('confirmation_ref'), 'font choice')
    font_dir = SKILL / 'assets/fonts/Pretendard'
    if 'font_files' in choice:
        inventory = choice['font_files']
        if not isinstance(inventory, dict) or set(inventory) != {'regular', 'bold'}:
            raise ValueError('Provided font_files require exact regular/bold OTF or TTF paths')
        fonts = [local_file(base, inventory[style]) for style in ('regular', 'bold')]
        faces = [inspect_font(path) for path in fonts]
        review = choice.get('font_license_review')
        if not isinstance(review, dict) or review.get('reviewed') is not True or review.get('cloud_use') != 'allowed':
            raise ValueError('Provided fonts require existing cloud-use license review')
        license_file = local_file(base, review.get('evidence_file'))
        if digest(license_file) != review.get('evidence_sha256'):
            raise ValueError('Font license evidence hash mismatch')
        if review.get('embedding') not in {'allowed', 'denied', 'unknown'}:
            raise ValueError('Record font embedding scope; adapter does not embed fonts')
        if any(face['sha256'] not in review.get('font_sha256', []) for face in faces):
            raise ValueError('Exact regular/bold font hashes must be covered by license review')
        review = {**review, 'evidence_file': str(license_file)}
    else:
        if family_name != 'Pretendard':
            raise ValueError('Selected font files unavailable; no silent fallback. Provide OTF/TTF plus license or explicitly choose bundled Pretendard')
        fonts = [font_dir / f'Pretendard-{style}.otf' for style in ('Regular', 'Bold')]
        faces = [inspect_font(path) for path in fonts]
        license_file = font_dir / 'LICENSE.txt'
        review = {'reviewed': True, 'cloud_use': 'allowed', 'embedding': 'allowed',
                  'evidence_file': str(license_file), 'evidence_sha256': digest(license_file),
                  'font_sha256': [face['sha256'] for face in faces]}
    for face, style in zip(faces, ('Regular', 'Bold')):
        if (family_name.casefold() not in [alias.casefold() for alias in face['family_aliases']]
                or style.casefold() not in [alias.casefold() for alias in face['style_aliases']]):
            raise ValueError('Provided regular/bold face identity differs from selected font; no substitution')
    brief = {'endpoint': 'cloud', 'font_policy': {'basis': 'specified', 'values': {'body': family_name}},
             'font_requests': [{'family': family_name, 'style': style, 'sha256': face['sha256']} for style, face in zip(('Regular', 'Bold'), faces)],
             'font_files': [str(path) for path in fonts], 'font_license_review': review}
    pages = plan.get('pages')
    if (not isinstance(pages, list) or not pages or type(plan.get('slide_count')) is not int
            or plan['slide_count'] != len(pages)):
        raise ValueError('Explicit slide_count must equal the complete nonempty page sequence')
    audit_pages = []
    for page in pages:
        if not isinstance(page, dict): raise ValueError('Every page must be an object')
        label(page.get('title'), 'page title')
        source = page.get('source', {})
        if not isinstance(source, dict): raise ValueError('Every source must be an object')
        if source.get('review_origin') not in {'local_model_visual_analysis', 'parent_model_visual_analysis'}:
            raise ValueError('Explicit local or parent-model review origin required')
        if not SHA.fullmatch(str(source.get('sha256', ''))):
            raise ValueError('Original photo SHA256 required')
        size = source.get('pixel_size')
        if not isinstance(size, list) or len(size) != 2 or any(type(v) is not int or v <= 0 for v in size):
            raise ValueError('Original photo pixel_size requires two positive integers')
        if source.get('input_type') not in IMAGE_TYPES:
            raise ValueError('Original source must be a supported photograph')
        label(source.get('visual_review_ref'), 'model visual review')
        if source.get('file'):
            actual = inspect_reference(local_file(base, source['file']))
            if (actual['input_sha256'] != source['sha256'] or actual['pixel_size'] != size
                    or actual['input_type'] != source['input_type']):
                raise ValueError('Actual source bytes differ from model source facts')
            local_verified = True
        else:
            if source.get('review_origin') != 'parent_model_visual_analysis':
                raise ValueError('Unavailable original bytes require explicit parent-model provenance')
            actual = {'input_type': source['input_type'], 'input_sha256': source['sha256'], 'pixel_size': size}
            local_verified = False
        bound = bind_photo_observations(actual, page.get('observations'))
        if not bound['reconstruction_contract']['reference_pixels_inspected']:
            raise ValueError('Model must inspect source pixels before reconstruction')
        observations = page['observations']
        if abs(observations['slide_aspect_ratio'] - CANVAS[0] / CANVAS[1]) > 0.03:
            raise ValueError('Confirmed slide canvas conflicts with model screen aspect estimate')
        items = {}
        for item in observations['items']:
            items.setdefault(item['element_id'], {})[item['property']] = item
        elements = page.get('elements')
        if not isinstance(elements, list) or not elements:
            raise ValueError('Model must provide the complete confirmed element sequence')
        seen, fitted = set(), {}
        native_count = 0
        for element in elements:
            if not isinstance(element, dict): raise ValueError('Every model element must be an object')
            eid = element.get('id')
            if not isinstance(eid, str) or not ID.fullmatch(eid) or eid in seen or eid not in items:
                raise ValueError('Unique safe element id bound to source observations required')
            seen.add(eid)
            kind = element.get('type')
            if kind not in {'rect', 'line', 'text', 'image'}:
                raise ValueError('Supported editable primitives: rect, line, text; supplied image crop only')
            if 'text' in items[eid] and kind != 'text':
                raise ValueError('Observed text must remain an editable text target; matching an id alone does not preserve it')
            supported = {'id', 'type', 'bounds_px', 'source_bounds_px'} | {
                'rect': {'color'}, 'line': {'color', 'direction', 'stroke_width_px'},
                'text': {'text', 'font_pt', 'font_family', 'font_weight', 'color', 'align', 'line_height_px', 'text_override'},
                'image': {'asset_file', 'asset_sha256'},
            }[kind]
            if set(element) - supported:
                raise ValueError('Unsupported element fields cannot be silently discarded: ' + ', '.join(sorted(set(element) - supported)))
            box(element.get('bounds_px'), CANVAS, 'rectified bounds_px')
            source_bounds = box(element.get('source_bounds_px'), size, 'source_bounds_px')
            if not any(item['source_bounds_px'] == source_bounds for item in items[eid].values()):
                raise ValueError('Element photo bounds differ from its bound observations')
            geometry = items[eid].get('geometry_px') or items[eid].get('image_crop_px')
            if not geometry or geometry['source_bounds_px'] != source_bounds:
                raise ValueError('Every element needs matching source geometry/crop provenance')
            expected_render = 'preserved_image_crop' if kind == 'image' else 'native_editable'
            if any(item['render_as'] != expected_render for item in items[eid].values()):
                raise ValueError('Element editability conflicts with observation targets')
            if kind != 'image': native_count += 1
            if kind == 'text':
                text = label(element.get('text'), 'text')
                if any(ord(c) < 32 and c != '\n' for c in text):
                    raise ValueError('Control characters are unsupported in literal text')
                if any(not row.strip() for row in text.split('\n')):
                    raise ValueError('Minimal adapter requires nonempty literal text lines; use separate positioned boxes')
                observed_text = items[eid].get('text')
                if not observed_text or observed_text['basis'] not in {'visible_text', 'user_instruction'}:
                    raise ValueError('Literal text requires visible or user-confirmed text provenance')
                override = element.get('text_override')
                if override is not None:
                    confirmed_text = isinstance(override, dict) and override.get('basis') == 'user_instruction' and override.get('certainty') == 'confirmed'
                    synthetic_text = (estimated_review and isinstance(override, dict)
                                      and override.get('basis') == 'user_authorized_synthetic' and override.get('certainty') == 'model_written')
                    if (not (confirmed_text or synthetic_text)
                            or override.get('original_observed_text') != observed_text['value']):
                        raise ValueError('Text replacement must preserve original observation and user confirmation')
                    label(override.get('source_ref'), 'text replacement')
                    if synthetic_text:
                        disclosure = label(page.get('synthetic_content_disclosure'), 'synthetic content disclosure')
                        if not any(e.get('type') == 'text' and disclosure in e.get('text', '') for e in elements):
                            raise ValueError('User-authorized synthetic text needs an on-slide disclosure')
                elif observed_text['value'] != text:
                    raise ValueError('Changed literal text requires explicit user-confirmed text_override')
                pt = number(element.get('font_pt'), 'font_pt')
                line = number(element.get('line_height_px'), 'line_height_px')
                if element.get('font_family') != choice['family'] or element.get('font_weight') not in (400, 700):
                    raise ValueError('Chosen font must match supported regular/bold face')
                if element.get('align') not in {'left', 'center', 'right'}:
                    raise ValueError('Explicit text align required')
                color(element.get('color'))
                for prop, value in [('font_size_pt', pt), ('line_spacing_pt', line * 72 / 96)]:
                    item = items[eid].get(prop)
                    if item and item['value'] is not None and abs(item['value'] - value) > 1e-6:
                        raise ValueError(f'{prop} differs from its source/user observation')
                for prop in ('paragraph_spacing_pt', 'tracking_pt'):
                    item = items[eid].get(prop)
                    if item and item['value'] not in (None, 0):
                        raise ValueError(f'Nonzero {prop} is unsupported; supply separately positioned literal boxes')
                for prop, value in [('alignment', element['align']), ('color', element['color'])]:
                    item = items[eid].get(prop)
                    if item and item['value'] is not None and str(item['value']).lower() != str(value).lower():
                        raise ValueError(f'{prop} differs from the recorded literal style')
                weight = items[eid].get('font_weight')
                if weight and weight['value'] is not None:
                    normalized = {'normal': 400, 'bold': 700}.get(weight['value'], weight['value'])
                    if normalized != element['font_weight']: raise ValueError('font_weight differs from literal style')
                family = items[eid].get('font_family')
                if family and family['basis'] == 'user_instruction' and family['value'] != element['font_family']:
                    raise ValueError('User-specified font cannot be replaced by the adapter design choice')
                font = fonts[1 if element['font_weight'] == 700 else 0]
                contract = {'role': 'body', 'text': text, 'font_file': str(font), 'font_size_pt': pt,
                            'line_spacing_pt': line * 72 / 96, 'paragraph_after_pt': 0, 'paragraph_before_pt': 0,
                            'letter_spacing_pt': 0, 'box_px': {'width': element['bounds_px'][2], 'height': element['bounds_px'][3]},
                            'margins_px': {'left': 0, 'right': 0, 'top': 0, 'bottom': 0}}
                fit = measure_text_fit(contract, base, brief, base)
                # Preserve literal lines/spaces; no implicit reflow or font shrink.
                from PIL import ImageFont
                face = ImageFont.truetype(str(font), round(pt * 96 / 72 * 4))
                literal_rows = text.split('\n')
                widths = [face.getlength(row) / 4 for row in literal_rows]
                baseline = face.getmetrics()[0] / 4
                last_ink = max(baseline + i * line + face.getbbox(row, anchor='ls')[3] / 4 for i, row in enumerate(literal_rows))
                if (fit['overflow'] or len(fit['lines']) != len(literal_rows)
                        or max(widths) > contract['box_px']['width'] or last_ink > contract['box_px']['height']):
                    raise ValueError(f'{eid}: text overflow; model must revise explicit lines/box/content, never shrink')
                fitted[eid] = {'measurement': fit, 'literal_widths_px': widths, 'baseline_px': baseline,
                               'literal_text_sha256': hashlib.sha256(text.encode()).hexdigest()}
            elif kind in {'rect', 'line'}:
                color(element.get('color'))
                if kind == 'line':
                    number(element.get('stroke_width_px'), 'stroke_width_px')
                    if element.get('direction') not in {'horizontal', 'vertical'}:
                        raise ValueError('Minimal adapter supports only horizontal/vertical lines')
                    if element['stroke_width_px'] > element['bounds_px'][3 if element['direction'] == 'horizontal' else 2]:
                        raise ValueError('Line stroke must fit inside the confirmed rectified bounds')
            else:
                asset = local_file(base, element.get('asset_file'))
                if asset.suffix.lower() not in IMAGE_TYPES or digest(asset) != element.get('asset_sha256'):
                    raise ValueError('Supplied image crop bytes/hash required')
                from PIL import Image
                with Image.open(asset) as picture:
                    picture.load(); asset_ratio = picture.width / picture.height
                _, _, w, h = element['bounds_px']
                if abs(asset_ratio / (w / h) - 1) > 0.01 or w * h >= CANVAS[0] * CANVAS[1] * 0.8:
                    raise ValueError('Image crop must preserve aspect and cannot replace the whole slide')
        omissions = page.get('omitted_decorations', [])
        if not isinstance(omissions, list): raise ValueError('Explicit omitted_decorations list required')
        omitted_ids = set()
        for omission in omissions:
            eid = omission.get('element_id')
            if eid not in items or eid in seen or eid in omitted_ids:
                raise ValueError('Omitted decoration must be a distinct observed element')
            if 'text' in items[eid]: raise ValueError('Required observed text cannot be silently omitted')
            label(omission.get('reason'), 'omission reason')
            if estimated_review and omission.get('basis') == 'visual_estimate':
                if omission.get('certainty') != 'estimated' or 'confirmation_ref' in omission:
                    raise ValueError('Model omission must remain estimated, not user-confirmed')
                label(omission.get('review_ref'), 'model omission review')
            else:
                if estimated_review and (omission.get('basis') != 'user_instruction' or omission.get('certainty') != 'confirmed'):
                    raise ValueError('Omission needs explicit model-estimate or actual user-instruction provenance')
                label(omission.get('confirmation_ref'), 'omission confirmation')
            omitted_ids.add(eid)
        if seen | omitted_ids != set(items) or not native_count:
            raise ValueError('Account for every observed element and retain native editable targets')
        audit_pages.append({'source': source, 'source_bytes_verified_locally': local_verified,
                            'observations': observations, 'reconstruction_contract': bound['reconstruction_contract'],
                            'elements': elements, 'omitted_decorations': omissions, 'text_fit': fitted})
    return {'schema_version': 1, 'validated': True, 'font_faces': faces, 'font_license_sha256': digest(license_file),
            'font_files': [str(p) for p in fonts], 'pages': audit_pages,
            'review_only': estimated_review, 'layout_basis': 'visual_estimate' if estimated_review else 'confirmed_input',
            'layout_review_ref': plan.get('layout_review_ref'), 'font_choice': choice,
            'individual_style_user_approval_verified': False,
            'model_analysis_performed_by_adapter': False, 'automatic_ocr_performed': False,
            'source_font_identity_verified': False, 'source_master_theme_recovered': False,
            'visual_source_comparison_pending': True, 'user_delivery_ready': False}


def make_svg(page: dict, audit: dict) -> str:
    """Translate confirmed primitives; the existing compiler owns PPTX writing."""
    root = ET.Element(f'{{{SVG}}}svg', {'width': '1280', 'height': '720', 'viewBox': '0 0 1280 720',
        'data-pptx-master': 'photo-review-master', 'data-pptx-master-name': 'Photo reconstruction review',
        'data-pptx-layout': 'photo-literal-fixed', 'data-pptx-layout-name': 'Photo literal fixed composition'})
    # Intentional new white Master + shared zero-slot Layout. All reconstructed
    # evidence is Slide-local and editable; no claim of source topology recovery.
    ET.SubElement(root, f'{{{SVG}}}rect', {'id': 'review-master-background', 'x': '0', 'y': '0',
        'width': '1280', 'height': '720', 'fill': '#FFFFFF', 'data-pptx-layer': 'master', 'data-pptx-editable': 'false'})
    for element in page['elements']:
        x, y, w, h = element['bounds_px']
        # Unmarked visible atoms are Slide-local in the established structured
        # compiler. Its explicit 'slide' mark is reserved for canvas backgrounds.
        attrs = {'id': element['id']}
        kind = element['type']
        if kind == 'rect':
            attrs.update({'x': str(x), 'y': str(y), 'width': str(w), 'height': str(h), 'fill': element['color']})
        elif kind == 'line':
            horizontal = element['direction'] == 'horizontal'
            attrs.update({'x1': str(x if horizontal else x + w / 2), 'y1': str(y + h / 2 if horizontal else y),
                          'x2': str(x + w if horizontal else x + w / 2), 'y2': str(y + h / 2 if horizontal else y + h),
                          'stroke': element['color'], 'stroke-width': str(element['stroke_width_px'])})
        elif kind == 'text':
            anchor = {'left': 'start', 'center': 'middle', 'right': 'end'}[element['align']]
            tx = x if anchor == 'start' else x + w / 2 if anchor == 'middle' else x + w
            baseline = y + audit['text_fit'][element['id']]['baseline_px']
            attrs.update({'x': str(tx), 'y': str(baseline), 'fill': element['color'],
                          'font-family': element['font_family'], 'font-size': str(element['font_pt'] * 96 / 72),
                          'font-weight': str(int(element['font_weight'])), 'text-anchor': anchor,
                          '{http://www.w3.org/XML/1998/namespace}space': 'preserve'})
            node = ET.SubElement(root, f'{{{SVG}}}text', attrs)
            rows = element['text'].split('\n')
            for i, row in enumerate(rows):
                ET.SubElement(node, f'{{{SVG}}}tspan', {'x': str(tx), 'y': str(baseline + i * element['line_height_px'])}).text = row
            continue
        else:
            attrs.update({'x': str(x), 'y': str(y), 'width': str(w), 'height': str(h),
                          'href': '../images/' + element['asset_sha256'] + local_file(Path('.'), element['asset_file']).suffix.lower(),
                          'preserveAspectRatio': 'xMidYMid meet'})
        ET.SubElement(root, f'{{{SVG}}}{kind}', attrs)
    ET.indent(root)
    return ET.tostring(root, encoding='unicode') + '\n'


def native_readback(path: Path, plan: dict) -> dict:
    from pptx import Presentation
    deck = Presentation(str(path))
    if len(deck.slides) != plan['slide_count']: raise ValueError('Native readback slide count mismatch')
    checks = []
    def leaves(shapes):
        for shape in shapes:
            if hasattr(shape, 'shapes'): yield from leaves(shape.shapes)
            else: yield shape
    for slide, page in zip(deck.slides, plan['pages']):
        shapes = list(leaves(slide.shapes))
        actual_texts = Counter(shape.text for shape in shapes if shape.has_text_frame and shape.text)
        # The SVG converter may use one native text box per explicit tspan line.
        expected_texts = Counter(row for element in page['elements'] if element['type'] == 'text' for row in element['text'].split('\n'))
        if actual_texts != expected_texts: raise ValueError('Literal native text readback mismatch')
        expected_points = Counter((row, round(element['font_pt'] * 100), element['font_weight'] == 700)
                                  for element in page['elements'] if element['type'] == 'text' for row in element['text'].split('\n'))
        points = Counter()
        for shape in shapes:
            if not shape.has_text_frame or not shape.text: continue
            runs = [run for para in shape.text_frame.paragraphs for run in para.runs if run.text]
            if not runs or any(run.font.size is None for run in runs):
                raise ValueError('Native font size missing from a literal text run')
            sizes = {round(run.font.size.pt * 100) for run in runs}
            if any(run.font.name != plan['font_choice']['family'] for para in shape.text_frame.paragraphs for run in para.runs):
                raise ValueError('Native font family differs from confirmed design choice')
            if len(sizes) != 1: raise ValueError('Native font size missing or mixed unexpectedly')
            weights = {run.font.bold is True for run in runs}
            if len(weights) != 1: raise ValueError('Native font weight mixed unexpectedly')
            points[(shape.text, next(iter(sizes)), next(iter(weights)))] += 1
        if points != expected_points: raise ValueError('Native font size or weight differs from the recorded plan')
        native = sum(not (shape.has_text_frame and shape.text) and shape.shape_type != 13 for shape in shapes)
        images = sum(shape.shape_type == 13 for shape in shapes)
        if native < sum(e['type'] in {'rect', 'line'} for e in page['elements']) or images != sum(e['type'] == 'image' for e in page['elements']):
            raise ValueError('Native shape/image editability readback mismatch')
        checks.append({'native_text_lines': sum(actual_texts.values()), 'native_shapes': native, 'images': images,
                       'literal_text_and_font_pt_verified': True})
    return {'slides': checks, 'slide_size_emu': [deck.slide_width, deck.slide_height], 'powerpoint_application_opened': False}


def runtime() -> dict:
    names = {'RUNTIME_NODE': 'CODEX_PRIMARY_RUNTIME_NODE', 'RUNTIME_NODE_MODULES': 'CODEX_PRIMARY_RUNTIME_NODE_MODULES',
             'RUNTIME_PYTHON': 'CODEX_PRIMARY_RUNTIME_PYTHON'}
    env = dict(os.environ)
    for target, source in names.items():
        value = env.get(source, '')
        if not Path(value).is_absolute() or not Path(value).exists():
            raise ValueError(f'Supplied {source} unavailable; no runtime fallback/installation')
        env[target] = value
    root = Path(env.get('CODEX_PRIMARY_RUNTIME', ''))
    if not root.is_absolute() or not (root / 'dependencies/bin/override').is_dir():
        raise ValueError('Supplied runtime bin directory unavailable')
    env['RUNTIME_BIN_DIR'] = str(root / 'dependencies/bin/override')
    return env


def build_review(plan_path: Path, parent: Path, shared_skill: Path) -> dict:
    plan_bytes = plan_path.read_bytes()
    plan = json.loads(plan_bytes)
    audit = validate_plan(plan, plan_path.parent)
    verify_fitted_font_files(audit)
    env = runtime()
    shared_skill = shared_skill.resolve()
    for helper in ['artifact_tool_utils.mjs', 'runtime_helpers.mjs', 'mark_artifact_operation_started.mjs',
                   'inspect_presentation_package_integrity.py', 'inspect_presentation_layout_geometry.py']:
        if not (shared_skill / 'container_tools' / helper).is_file():
            raise ValueError('Shared Presentations helper unavailable: ' + helper)
    parent = parent.resolve()
    if not parent.is_relative_to(REPO / 'projects'):
        raise ValueError('Photo review output must stay in gitignored repository projects/')
    workspace = Path(ProjectManager(parent).init_project(plan['project_name'], 'ppt169'))
    write_project_gate(workspace, make_exempt_record('create-template'))
    templates = workspace / 'templates'; templates.mkdir()
    analysis = workspace / 'analysis'
    write_json(analysis / 'recorded-model-plan.json', plan)
    (analysis / 'received-model-plan.json').write_bytes(plan_bytes)
    audit['plan_file_sha256'] = hashlib.sha256(plan_bytes).hexdigest()
    audit['plan_snapshot_file'] = 'received-model-plan.json'
    audit['plan_file'] = str(plan_path)
    audit['plan_relative_paths_base'] = str(plan_path.parent)
    audit['source_files_sha256'] = {path.name: digest(path) for path in [Path(__file__), SCRIPTS / 'reference_intake.py', SCRIPTS / 'template_preview_pptx.py']}
    write_json(analysis / 'reconstruction-audit.json', audit)
    def run(args: list[str], name: str) -> None:
        result = subprocess.run(args, cwd=REPO, env=env, text=True, capture_output=True)
        (analysis / (name + '.log')).write_text(result.stdout + result.stderr, encoding='utf-8')
        if result.returncode: raise ValueError(f'{name} failed ({result.returncode}); inspect {analysis / (name + ".log")}')
    roster = []
    for index, (page, page_audit) in enumerate(zip(plan['pages'], audit['pages']), 1):
        for element in page['elements']:
            if element['type'] == 'image':
                asset = local_file(plan_path.parent, element['asset_file'])
                destination = workspace / 'images' / (element['asset_sha256'] + asset.suffix.lower())
                asset_bytes = asset.read_bytes()
                if hashlib.sha256(asset_bytes).hexdigest() != element['asset_sha256']:
                    raise ValueError('Image asset changed after validation; original bytes must be preserved')
                if destination.exists():
                    if digest(destination) != element['asset_sha256']:
                        raise ValueError('Existing image asset differs from its declared source hash')
                else:
                    destination.write_bytes(asset_bytes)
        name = f'{index:03}_content'
        (templates / (name + '.svg')).write_text(make_svg(page, page_audit), encoding='utf-8')
        roster.append(f'| `{name}.svg` | Photo review page {index}; fixed editable composition |')
        if index == 1:
            run([env['RUNTIME_PYTHON'], str(SCRIPTS / 'svg_quality_checker.py'), str(templates / (name + '.svg')), '--template-mode', '--format', 'ppt169'], 'page-1-svg-quality')
    placeholders = '\n'.join(f'  {i:03}_content: []' for i in range(1, len(roster) + 1))
    (templates / 'design_spec.md').write_text(f'''---
deck_id: {plan['project_name']}
kind: deck
native_structure_mode: structured
canvas_format: ppt169
page_count: {len(roster)}
page_types: [content]
placeholders:
{placeholders}
---

# Photo reference review

## I. Template Overview
Private standard/literal photo reconstruction review. Input layout basis: {audit['layout_basis']}. No source Master/theme recovery.

## II. Color Scheme
Colors are explicit model/user inputs per element, retained in the reconstruction audit.

## III. Signature Design Elements
Rectified element bounds from the recorded model plan; estimates remain estimates. New white Master and one intentionally zero-slot Layout; all reconstructed elements are editable Slide-local objects. Literal text is not a reusable placeholder. This review is consumed by its explicit workspace path and is not registered globally.

## V. Page Roster
| File | Use |
| --- | --- |
{chr(10).join(roster)}
''', encoding='utf-8')
    run([env['RUNTIME_PYTHON'], str(SCRIPTS / 'svg_quality_checker.py'), str(templates), '--template-mode', '--format', 'ppt169'], 'all-svg-quality')
    run([env['RUNTIME_NODE'], str(shared_skill / 'container_tools/mark_artifact_operation_started.mjs'), '--operation-kind', 'create', '--expected-output-count', '1', '--output-format', 'pptx'], 'artifact-operation-marker')
    candidate = workspace / 'exports/review-candidate.pptx'
    run([env['RUNTIME_PYTHON'], str(SCRIPTS / 'template_preview_pptx.py'), str(workspace), '-o', str(candidate)], 'structured-compiler')
    readback = native_readback(candidate, plan)
    write_json(analysis / 'native-readback.json', readback)
    final = workspace / 'exports/review-final.pptx'
    profile = analysis / 'render-profile.json'
    write_json(profile, {'font_files': audit['font_files'], 'font_family': plan['font_choice']['family'], 'slide_count': plan['slide_count']})
    verify_fitted_font_files(audit)
    run([env['RUNTIME_NODE'], str(SCRIPTS / 'photo_review_finalize.mjs'), str(workspace), str(candidate), str(final), str(profile), str(shared_skill), ','.join(map(str, readback['slide_size_emu']))], 'shared-finalizer')
    verify_fitted_font_files(audit)
    run([env['RUNTIME_NODE'], str(SKILL / 'examples/korean_business/render_review.mjs'), str(final), str(workspace / 'rendered'), str(profile), str(shared_skill)], 'shared-all-page-render')
    render_manifest = workspace / 'rendered/render-manifest.json'
    font_binding = verify_render_font_binding(audit, json.loads(render_manifest.read_bytes()))
    font_binding_file = analysis / 'font-render-binding.json'
    write_json(font_binding_file, font_binding)
    final_readback = native_readback(final, plan)
    receipt = {'schema_version': 1, 'status': 'review_ready_visual_comparison_pending', 'workspace': str(workspace),
               'final_pptx': str(final), 'final_sha256': digest(final), 'slide_count': plan['slide_count'],
               'render_manifest': str(render_manifest), 'native_readback': final_readback,
               'font_render_binding': str(font_binding_file), 'fit_render_font_hashes_verified': True,
               'source_bytes_verified_locally': [page['source_bytes_verified_locally'] for page in audit['pages']],
               'review_only': audit['review_only'], 'layout_basis': audit['layout_basis'],
               'individual_style_user_approval_verified': False,
               'visual_source_comparison_pending': True, 'user_delivery_ready': False,
               'automatic_ocr_performed': False, 'source_font_identity_verified': False,
               'registered_active_template': False, 'owner': 'project-private create-template structured review'}
    write_json(analysis / 'orchestration-receipt.json', receipt)
    return receipt


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('plan', type=Path); parser.add_argument('--check', action='store_true')
    parser.add_argument('--workspace-parent', type=Path, default=REPO / 'projects')
    parser.add_argument('--presentations-skill-dir', type=Path)
    args = parser.parse_args(argv)
    try:
        if args.check:
            audit = validate_plan(json.loads(args.plan.read_text(encoding='utf-8')), args.plan.resolve().parent)
            result = {'validated': True, 'slide_count': len(audit['pages']), 'generation_performed': False,
                      'review_only': audit['review_only'], 'layout_basis': audit['layout_basis'],
                      'individual_style_user_approval_verified': False,
                      'user_delivery_ready': False, 'source_bytes_verified_locally': [p['source_bytes_verified_locally'] for p in audit['pages']]}
        else:
            if not args.presentations_skill_dir: raise ValueError('Read shared Presentations skill and provide its verified directory')
            result = build_review(args.plan.resolve(), args.workspace_parent, args.presentations_skill_dir)
        print(json.dumps(result, ensure_ascii=False, indent=2)); return 0
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print('Photo reconstruction blocked: ' + str(exc), file=sys.stderr); return 1


if __name__ == '__main__':
    raise SystemExit(main())
